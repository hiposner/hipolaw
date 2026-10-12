"""Cross-platform regression checks (Word COM integration must run on Windows)."""
import importlib.util
import sys
import unittest
from pathlib import Path
from unittest.mock import patch, Mock

APP=Path(__file__).resolve().parents[1]
SCRIPT=APP/'generate_legal_tables_v16_4_2.py'
spec=importlib.util.spec_from_file_location('hipolaw_v163',SCRIPT)
mod=importlib.util.module_from_spec(spec)
sys.modules[spec.name]=mod
spec.loader.exec_module(mod)


class TestV163GridSetup(unittest.TestCase):
    def test_backward_compatibility_default_no(self):
        self.assertFalse(mod.read_whole_motion_grid_option({'pleading grid':'Auto'}))

    def test_yes_and_no(self):
        self.assertTrue(mod.read_whole_motion_grid_option({mod.WHOLE_MOTION_GRID_KEY:'Yes'}))
        self.assertFalse(mod.read_whole_motion_grid_option({mod.WHOLE_MOTION_GRID_KEY:'No'}))

    def test_ambiguous_answers_are_blocked(self):
        for x in ('Auto','1','Maybe','', 'true', '0'):
            with self.assertRaisesRegex(ValueError,'Yes or No'):
                mod.read_whole_motion_grid_option({mod.WHOLE_MOTION_GRID_KEY:x})

    def test_v16_4_excel_template(self):
        import openpyxl
        p=APP/'05_Review_Template_V16_4.xlsx'
        w=openpyxl.load_workbook(p,read_only=False)
        try:
            self.assertEqual(w.sheetnames,['Authorities','Contents','Setup','Aliases'])
            self.assertEqual(w['Setup']['A8'].value,'Align entire motion to 28-line pleading grid')
            self.assertEqual(w['Setup']['B8'].value,'No')
            self.assertEqual(w['Setup']['A9'].value,'HipoLaw document SHA256')
            self.assertEqual(w['Setup']['A10'].value,'HipoLaw document filename')
        finally:w.close()

    def test_real_lamotte_preflight_legacy(self):
        source=Path('tests/fixtures/optional_motion.docx')
        workbook=Path('tests/fixtures/optional_approved.xlsx')
        if not source.is_file() or not workbook.is_file():self.skipTest('sample source not present')
        args=mod.preflight(source,workbook)
        self.assertEqual(len(args[0]),37)
        self.assertEqual(len(args[1]),14)


class FakeFormat:
    def __init__(self):
        self.SpaceBefore=10
        self.SpaceAfter=11
        self.LineSpacingRule=1
        self.LineSpacing=18

class FakeRange:
    def __init__(self,pos,in_table=False):
        self.Start=pos
        self.ParagraphFormat=FakeFormat()
        self._table=in_table
    def Information(self,n):
        assert n==12
        return self._table

class FakeP:
    def __init__(self,pos,in_table=False):
        self.Range=FakeRange(pos,in_table)

class FakeDocument:
    def __init__(self):
        self.caption=FakeP(30)
        self.body=[FakeP(100),FakeP(200),FakeP(300,True),FakeP(400)]
        self.Content=type('Content',(),{'End':700})()
        self.sent_range=[]
    def Range(self,start,end):
        self.sent_range.append((start,end))
        return type('Body',(),{'Paragraphs':self.body})()

class TestWholeMotionCOMFormatting(unittest.TestCase):
    def test_no_leaves_document_untouched(self):
        d=FakeDocument()
        self.assertEqual(mod.apply_whole_motion_grid(d,'INTRODUCTION',enabled=False,
                  has_28_line_header=True,spacing_pt=24,log=[]),0)
        self.assertEqual(d.body[0].Range.ParagraphFormat.SpaceBefore,10)

    def test_yes_formats_body_but_not_caption_or_table(self):
        d=FakeDocument();log=[]
        with patch.object(mod,'heading_ranges',return_value=[type('R',(),{'Start':100})()]):
            n=mod.apply_whole_motion_grid(d,'INTRODUCTION',enabled=True,
                  has_28_line_header=True,spacing_pt=24,log=log)
        self.assertEqual(n,3)
        self.assertEqual(d.sent_range,[(100,700)])
        self.assertEqual(d.caption.Range.ParagraphFormat.SpaceBefore,10)
        self.assertEqual(d.body[2].Range.ParagraphFormat.SpaceBefore,10)
        for p in [d.body[0],d.body[1],d.body[3]]:
            fmt=p.Range.ParagraphFormat
            self.assertEqual((fmt.SpaceBefore,fmt.SpaceAfter,fmt.LineSpacingRule,fmt.LineSpacing),(0,0,4,24))
        self.assertIn('Footnotes', '\n'.join(log))

    def test_no_source_grid_stops(self):
        with self.assertRaisesRegex(ValueError,'28-line'):
            mod.apply_whole_motion_grid(FakeDocument(),'INTRODUCTION',enabled=True,
                          has_28_line_header=False,spacing_pt=24,log=[])


class TestPostPagination(unittest.TestCase):
    def test_mismatched_authority_page_blocks_success(self):
        with self.assertRaisesRegex(RuntimeError,'PAGINATION MISMATCH'):
            mod.ensure_pagination_matches(['REVIEW: Case X | TOA=[7] | body=[9]'])

    def test_cosmetic_grid_review_does_not_misclassify_pagination(self):
        mod.ensure_pagination_matches(['REVIEW: PDF p.3 entry 6.0pt off grid: Something'])

    def test_missing_toc_entry_blocks(self):
        with self.assertRaisesRegex(RuntimeError,'PAGINATION MISMATCH'):
            mod.ensure_pagination_matches(['REVIEW: missing or altered TOC label: I INTRODUCTION'])

    def test_word_pdfa_and_style_settings_untouched(self):
        import json
        data=json.loads((APP/'HipoLaw_Style_Settings.json').read_text())
        self.assertTrue(mod._validate_style(data))
        d=Mock()
        mod.export_court_pdfa(d,Path('C:/court.pdf'))
        self.assertTrue(d.ExportAsFixedFormat.call_args.kwargs['UseISO19005_1'])
        self.assertFalse(d.ExportAsFixedFormat.call_args.kwargs['IncludeDocProps'])


if __name__=='__main__':unittest.main(verbosity=2)

class HeaderGridTests(unittest.TestCase):
    class Paragraph:
        def __init__(self,text):
            self.Range=type('R',(),{})()
            self.Range.Text=text
            self.Range.ParagraphFormat=FakeFormat()
    class Header:
        def __init__(self):
            self.LinkToPrevious=True
            self.Range=type('HeaderRange',(),{})()
            self.Range.Paragraphs=[HeaderGridTests.Paragraph(' '.join(map(str,range(1,29))))]
            self.Range.Tables=[]
    class Section:
        def __init__(self):self.header=HeaderGridTests.Header()
        def Headers(self,which):return self.header
    class Doc:
        def __init__(self):self.sec=HeaderGridTests.Section()
        def Sections(self,index):return self.sec

    def test_header_number_baseline_shift_only_when_yes(self):
        d=self.Doc();log=[]
        self.assertEqual(mod.align_generated_body_header_numbers(d,3,enabled=True,shift_pt=9.5,log=log),1)
        self.assertFalse(d.sec.header.LinkToPrevious)
        self.assertEqual(d.sec.header.Range.Paragraphs[0].Range.ParagraphFormat.SpaceBefore,19.5)
        self.assertIn('body-section', '\n'.join(log))

    def test_header_skipped_when_no(self):
        d=self.Doc()
        self.assertEqual(mod.align_generated_body_header_numbers(d,3,enabled=False,shift_pt=9.5,log=[]),0)
        self.assertTrue(d.sec.header.LinkToPrevious)

    def test_unrecognized_header_fails_safely(self):
        d=self.Doc()
        d.sec.header.Range.Paragraphs[0].Range.Text='This header has no numbering.'
        with self.assertRaisesRegex(RuntimeError,'28 printed pleading numbers'):
            mod.align_generated_body_header_numbers(d,3,enabled=True,shift_pt=9.5,log=[])

    def test_baseline_setting_can_be_changed_without_rebuild(self):
        import json
        a=json.loads((APP/'HipoLaw_Style_Settings.json').read_text())
        self.assertEqual(a['pleading_grid']['body_header_shift_pt'],9.5)
        mod._validate_style(a)
        a['pleading_grid']['body_header_shift_pt']=6
        mod._validate_style(a)
        a['pleading_grid']['body_header_shift_pt']=22
        with self.assertRaisesRegex(ValueError,'body_header_shift_pt'):
            mod._validate_style(a)
